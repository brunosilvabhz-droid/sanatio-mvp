from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.clinical import ProdutoAntimicrobiano


def clean(value) -> str:
    return str(value if value is not None else "").strip()


def read_products(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]
    required = {"CD_PRODUTO", "DS_PRODUTO", "QT", "UNIDADE", "CODIGO_ATC"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Colunas obrigatorias ausentes: {', '.join(sorted(missing))}")
    products = []
    for values in rows:
        row = {headers[index]: clean(value) for index, value in enumerate(values)}
        if not row["CD_PRODUTO"]:
            continue
        quantity = row["QT"].replace(",", ".")
        if row["UNIDADE"] == "UI" and "." in quantity:
            quantity = quantity.replace(".", "")
        products.append(
            {
                "codigo_produto": row["CD_PRODUTO"],
                "descricao": row["DS_PRODUTO"],
                "quantidade": quantity,
                "unidade_medida": row["UNIDADE"],
                "codigo_atc": row["CODIGO_ATC"].upper(),
            }
        )
    workbook.close()
    if len({item["codigo_produto"] for item in products}) != len(products):
        raise ValueError("A planilha possui codigos de produto duplicados")
    return products


def main() -> None:
    parser = argparse.ArgumentParser(description="Importa o catalogo de produtos antimicrobianos")
    parser.add_argument("xlsx_path", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    products = read_products(args.xlsx_path)
    if not args.apply:
        print({"mode": "DRY_RUN", "products": len(products)})
        return
    with SessionLocal() as db:
        existing = {item.codigo_produto: item for item in db.scalars(select(ProdutoAntimicrobiano))}
        created = updated = 0
        for data in products:
            item = existing.get(data["codigo_produto"])
            if not item:
                item = ProdutoAntimicrobiano(codigo_produto=data["codigo_produto"])
                db.add(item)
                created += 1
            else:
                updated += 1
            for key, value in data.items():
                setattr(item, key, value)
        db.commit()
    print({"mode": "APPLY", "products": len(products), "created": created, "updated": updated})


if __name__ == "__main__":
    main()
