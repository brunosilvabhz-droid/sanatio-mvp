import csv
import io
import json
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, require_admin
from app.core.database import get_db
from app.models.epidemiology_reference import (
    FechamentoAnvisa,
    ImportacaoReferenciaEpidemiologica,
    LogAuditoria,
    PerfilEstabelecimento,
    ReferenciaEpidemiologica,
)
from app.models.user import User
from app.schemas.epidemiology_reference import (
    BenchmarkComparacaoRead,
    FechamentoAnvisaRead,
    FechamentoStatusUpdate,
    ImportacaoReferenciaRead,
    PerfilEstabelecimentoRead,
    PerfilEstabelecimentoUpdate,
    ReferenciaEpidemiologicaCreate,
    ReferenciaEpidemiologicaRead,
)
from app.services.benchmark_epidemiologico_service import INDICADORES, BenchmarkEpidemiologicoService
from app.services.servico_fontes_publicas import CNESAdapter

router = APIRouter(prefix="/epidemiology/public-references", tags=["Referencias epidemiologicas"])

CSV_COLUMNS = [
    "codigo_indicador",
    "nome_indicador",
    "tipo_unidade",
    "populacao_referencia",
    "regiao",
    "uf",
    "ano_referencia",
    "p10",
    "p25",
    "p50",
    "p75",
    "p90",
    "unidade_medida",
    "fonte",
    "url_fonte",
]

REQUIRED_IMPORT_COLUMNS = [
    "codigo_indicador",
    "nome_indicador",
    "tipo_unidade",
    "populacao_referencia",
    "ano_referencia",
    "unidade_medida",
    "fonte",
]


def _audit(db: Session, user: User | None, acao: str, registro: str, before: object | None = None, after: object | None = None) -> None:
    db.add(
        LogAuditoria(
            usuario_id=user.id if user else None,
            acao=acao,
            registro=registro,
            valor_anterior=json.dumps(before, ensure_ascii=False, default=str) if before is not None else None,
            valor_novo=json.dumps(after, ensure_ascii=False, default=str) if after is not None else None,
        )
    )


def _to_float(value: str | None) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    return float(str(value).replace(",", "."))


def _to_int(value: str | None, field: str) -> int:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field} é obrigatório")
    return int(value)


def _today_period() -> str:
    today = date.today()
    return f"{today.year:04d}-{today.month:02d}"


def _validate_percentiles(values: list[float | None]) -> None:
    previous = None
    for value in values:
        if value is None:
            continue
        if previous is not None and value < previous:
            raise ValueError("Percentis devem seguir a ordem P10 <= P25 <= P50 <= P75 <= P90")
        previous = value


def _normalize_import_row(row: dict) -> dict[str, str]:
    return {
        str(key or "").strip().lower(): str(value or "").strip()
        for key, value in row.items()
        if key is not None
    }


def _read_csv(raw_content: bytes) -> list[dict[str, str]]:
    try:
        content = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError:
        content = raw_content.decode("cp1252")
    if not content.strip():
        return []
    sample = content[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        reader = csv.DictReader(io.StringIO(content), dialect=dialect)
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.DictReader(io.StringIO(content), delimiter=delimiter)
    return [_normalize_import_row(row) for row in reader if any(str(value or "").strip() for value in row.values())]


@router.get("/references", response_model=list[ReferenciaEpidemiologicaRead])
def references(
    indicador: str | None = None,
    ano: int | None = None,
    tipo_unidade: str | None = None,
    ativo: bool | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> list[ReferenciaEpidemiologica]:
    query = select(ReferenciaEpidemiologica)
    if indicador:
        query = query.where((ReferenciaEpidemiologica.codigo_indicador == indicador) | (ReferenciaEpidemiologica.nome_indicador.ilike(f"%{indicador}%")))
    if ano:
        query = query.where(ReferenciaEpidemiologica.ano_referencia == ano)
    if tipo_unidade:
        query = query.where(ReferenciaEpidemiologica.tipo_unidade == tipo_unidade)
    if ativo is not None:
        query = query.where(ReferenciaEpidemiologica.ativo.is_(ativo))
    return list(db.scalars(query.order_by(ReferenciaEpidemiologica.codigo_indicador, ReferenciaEpidemiologica.ano_referencia.desc())))


@router.post("/references", response_model=ReferenciaEpidemiologicaRead, dependencies=[Depends(require_admin)])
def create_reference(payload: ReferenciaEpidemiologicaCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> ReferenciaEpidemiologica:
    item = ReferenciaEpidemiologica(**payload.model_dump())
    db.add(item)
    _audit(db, user, "referencia_epidemiologica_criada", payload.codigo_indicador, after=payload.model_dump())
    db.commit()
    db.refresh(item)
    return item


@router.patch("/references/{reference_id}/active", response_model=ReferenciaEpidemiologicaRead, dependencies=[Depends(require_admin)])
def set_reference_active(reference_id: int, active: bool, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> ReferenciaEpidemiologica:
    item = db.get(ReferenciaEpidemiologica, reference_id)
    if not item:
        raise HTTPException(status_code=404, detail="Referência não encontrada")
    before = {"ativo": item.ativo}
    item.ativo = active
    _audit(db, user, "referencia_epidemiologica_status", f"referencia:{reference_id}", before=before, after={"ativo": active})
    db.commit()
    db.refresh(item)
    return item


@router.post("/references/import-csv", response_model=ImportacaoReferenciaRead, dependencies=[Depends(require_admin)])
async def import_csv(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> ImportacaoReferenciaEpidemiologica:
    original_filename = file.filename or "referencias_epidemiologicas"
    filename = original_filename.lower()
    raw_content = await file.read()
    if filename.endswith(".csv"):
        rows = _read_csv(raw_content)
    elif filename.endswith(".xlsx"):
        rows = _read_xlsx(raw_content)
    else:
        raise HTTPException(status_code=422, detail="Importação aceita arquivos CSV ou Excel .xlsx")
    if not rows:
        raise HTTPException(status_code=422, detail="O arquivo está vazio ou não possui linhas de dados")
    missing = [column for column in REQUIRED_IMPORT_COLUMNS if column not in rows[0]]
    if missing:
        raise HTTPException(status_code=422, detail=f"Arquivo sem colunas obrigatórias: {', '.join(missing)}")

    importacao = ImportacaoReferenciaEpidemiologica(nome_arquivo=original_filename, fonte="CSV/Excel", usuario_id=user.id, status="EM_PROCESSAMENTO")
    db.add(importacao)
    db.flush()
    errors: list[str] = []
    inserted = 0
    version = datetime.now(timezone.utc).strftime("csv-%Y%m%d%H%M%S")
    seen_keys: set[tuple] = set()
    for index, row in enumerate(rows, start=2):
        try:
            codigo = (row.get("codigo_indicador") or "").strip().upper()
            nome = (row.get("nome_indicador") or "").strip()
            tipo = (row.get("tipo_unidade") or "").strip().upper()
            populacao = (row.get("populacao_referencia") or "").strip().upper()
            fonte = (row.get("fonte") or "").strip()
            ano = _to_int(row.get("ano_referencia"), "ano_referencia")
            if not codigo or not nome or not tipo or not populacao or not fonte:
                raise ValueError("codigo_indicador, nome_indicador, tipo_unidade, populacao_referencia e fonte são obrigatórios")
            key = (codigo, tipo, populacao, row.get("regiao") or None, row.get("uf") or None, ano, row.get("periodo_referencia") or None)
            if key in seen_keys:
                raise ValueError("Referência duplicada no arquivo")
            seen_keys.add(key)
            p10 = _to_float(row.get("p10"))
            p25 = _to_float(row.get("p25"))
            p50 = _to_float(row.get("p50"))
            p75 = _to_float(row.get("p75"))
            p90 = _to_float(row.get("p90"))
            _validate_percentiles([p10, p25, p50, p75, p90])
            item = ReferenciaEpidemiologica(
                codigo_indicador=codigo,
                nome_indicador=nome,
                tipo_unidade=tipo,
                populacao_referencia=populacao,
                regiao=(row.get("regiao") or None),
                uf=(row.get("uf") or None),
                ano_referencia=ano,
                periodo_referencia=row.get("periodo_referencia") or None,
                p10=p10,
                p25=p25,
                p50=p50,
                p75=p75,
                p90=p90,
                unidade_medida=(row.get("unidade_medida") or "").strip() or "não informado",
                fonte=fonte,
                url_fonte=row.get("url_fonte") or None,
                versao_referencia=version,
                ativo=True,
            )
            db.add(item)
            inserted += 1
            importacao.ano_referencia = ano
        except Exception as exc:
            errors.append(f"Linha {index}: {exc}")
    importacao.quantidade_registros = inserted
    importacao.quantidade_erros = len(errors)
    importacao.status = "CONCLUIDO_COM_ERROS" if errors else "CONCLUIDO"
    importacao.data_hora_fim = datetime.now(timezone.utc)
    importacao.mensagem_erro = "\n".join(errors[:20]) if errors else None
    _audit(db, user, "importacao_referencias_epidemiologicas", original_filename, after={"registros": inserted, "erros": len(errors), "versao": version})
    db.commit()
    db.refresh(importacao)
    return importacao


def _read_xlsx(raw_content: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="Dependência openpyxl não instalada no backend") from exc
    workbook = load_workbook(io.BytesIO(raw_content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    parsed = []
    for row in rows[1:]:
        parsed.append(_normalize_import_row({headers[index]: "" if value is None else str(value) for index, value in enumerate(row) if index < len(headers)}))
    return parsed


@router.get("/imports", response_model=list[ImportacaoReferenciaRead])
def imports(db: Session = Depends(get_db), _: User = Depends(get_current_user)) -> list[ImportacaoReferenciaEpidemiologica]:
    return list(db.scalars(select(ImportacaoReferenciaEpidemiologica).order_by(ImportacaoReferenciaEpidemiologica.data_hora_inicio.desc()).limit(50)))


@router.get("/benchmark/comparisons", response_model=list[BenchmarkComparacaoRead])
def benchmark_comparisons(
    periodo: str = Query(default_factory=_today_period),
    tipo_unidade: str = "UTI_ADULTO",
    indicador: str | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> list[dict]:
    service = BenchmarkEpidemiologicoService(db)
    specs = [spec for spec in INDICADORES if not indicador or indicador.upper() in {spec.codigo.upper(), spec.nome.upper()}]
    return [service.comparar(spec.codigo, periodo, tipo_unidade) for spec in specs]


@router.get("/establishment-profile", response_model=PerfilEstabelecimentoRead | None)
def establishment_profile(db: Session = Depends(get_db), _: User = Depends(get_current_user)) -> PerfilEstabelecimento | None:
    return db.scalar(select(PerfilEstabelecimento).order_by(PerfilEstabelecimento.id))


@router.put("/establishment-profile", response_model=PerfilEstabelecimentoRead, dependencies=[Depends(require_admin)])
def update_establishment_profile(payload: PerfilEstabelecimentoUpdate, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> PerfilEstabelecimento:
    item = db.scalar(select(PerfilEstabelecimento).where(PerfilEstabelecimento.cnes == payload.cnes)) or PerfilEstabelecimento(cnes=payload.cnes)
    before = {"cnes": item.cnes, "nome_estabelecimento": item.nome_estabelecimento, "municipio": item.municipio, "uf": item.uf}
    for key, value in payload.model_dump().items():
        setattr(item, key, value)
    db.add(item)
    _audit(db, user, "perfil_estabelecimento_alterado", payload.cnes, before=before, after=payload.model_dump())
    db.commit()
    db.refresh(item)
    return item


@router.post("/establishment-profile/update-public-data", response_model=PerfilEstabelecimentoRead, dependencies=[Depends(require_admin)])
def update_public_establishment_data(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> PerfilEstabelecimento:
    item = db.scalar(select(PerfilEstabelecimento).order_by(PerfilEstabelecimento.id))
    if not item:
        raise HTTPException(status_code=422, detail="Cadastre o CNES antes de atualizar dados públicos")
    before = {"nome_estabelecimento": item.nome_estabelecimento, "erro_ultima_atualizacao": item.erro_ultima_atualizacao}
    try:
        data = CNESAdapter().consultar_estabelecimento(item.cnes)
        item.nome_estabelecimento = data.nome_estabelecimento or item.nome_estabelecimento
        item.municipio = data.municipio or item.municipio
        item.uf = data.uf or item.uf
        item.quantidade_leitos = data.quantidade_leitos if data.quantidade_leitos is not None else item.quantidade_leitos
        item.quantidade_leitos_uti = data.quantidade_leitos_uti if data.quantidade_leitos_uti is not None else item.quantidade_leitos_uti
        item.data_ultima_atualizacao_cnes = datetime.now(timezone.utc)
        item.fonte = data.fonte
        item.erro_ultima_atualizacao = None
    except Exception as exc:
        item.erro_ultima_atualizacao = str(exc)
    _audit(db, user, "perfil_estabelecimento_atualizacao_cnes", item.cnes, before=before, after={"erro": item.erro_ultima_atualizacao})
    db.commit()
    db.refresh(item)
    return item


@router.get("/anvisa-closure", response_model=FechamentoAnvisaRead)
def anvisa_closure(periodo: str = Query(default_factory=_today_period), tipo_unidade: str = "UTI_ADULTO", db: Session = Depends(get_db), _: User = Depends(get_current_user)) -> FechamentoAnvisa:
    item = db.scalar(select(FechamentoAnvisa).where(FechamentoAnvisa.periodo == periodo, FechamentoAnvisa.tipo_unidade == tipo_unidade))
    calculated = BenchmarkEpidemiologicoService(db).fechamento_anvisa(periodo, tipo_unidade)
    if not item:
        item = FechamentoAnvisa(**calculated)
        db.add(item)
    elif item.status != "validado":
        for key, value in calculated.items():
            setattr(item, key, value)
    db.commit()
    db.refresh(item)
    return item


@router.patch("/anvisa-closure/{closure_id}/status", response_model=FechamentoAnvisaRead)
def update_closure_status(closure_id: int, payload: FechamentoStatusUpdate, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> FechamentoAnvisa:
    if payload.status not in {"rascunho", "em_conferencia", "validado"}:
        raise HTTPException(status_code=422, detail="Status inválido")
    item = db.get(FechamentoAnvisa, closure_id)
    if not item:
        raise HTTPException(status_code=404, detail="Fechamento não encontrado")
    before = {"status": item.status}
    item.status = payload.status
    if payload.status == "validado":
        item.validado_por = user.id
        item.data_hora_validacao = datetime.now(timezone.utc)
    _audit(db, user, "fechamento_anvisa_status", f"fechamento:{closure_id}", before=before, after={"status": payload.status})
    db.commit()
    db.refresh(item)
    return item
